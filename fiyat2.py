# -*- coding: utf-8 -*-
import re
import shutil
import pandas as pd
from openpyxl import load_workbook

# ---------------- AYARLAR ----------------
POS_DOSYA = "Ürünler_Dışa_Aktar_-__BenimPOS.xlsx"
# DİKKAT: Aşağıya en son resimlerin çekildiği, şu an elindeki dosyanın tam adını yaz:
RESIMLI_DOSYA = "HARUN_CALISMA_RESIMLI_SON.xlsx" 
CIKTI_DOSYA = "HARUN_CALISMA_EN_SON_FINAL.xlsx"
# -----------------------------------------

def normalize(barkod):
    rakamlar = re.sub(r"\D", "", str(barkod))
    return rakamlar.lstrip("0")

# 1. BenimPOS dosyasından Fiyat 2 verilerini topla
print("BenimPOS dosyasından 'Fiyat 2' verileri okunuyor...")
pos = pd.read_excel(POS_DOSYA, sheet_name="Sheet1", header=1, dtype=str)
fiyat2_sozlugu = {}
for barkod, f2 in zip(pos["Ürün barkodu"], pos["Fiyat 2"]):
    anahtar = normalize(barkod)
    if anahtar and pd.notna(f2):
        fiyat2_sozlugu[anahtar] = f2

# 2. Resimli dosyayı kopyala ve aç (Resimlerin bozulmaması için çok önemli)
print("Resimli dosya yükleniyor, lütfen bekleyin...")
shutil.copy(RESIMLI_DOSYA, CIKTI_DOSYA)
wb = load_workbook(CIKTI_DOSYA)
ws = wb["Ürünlerinizi Burada Listeleyin"]

# Başlıkları bul
basliklar = {str(c.value).strip(): c.column for c in ws[1] if c.value}
col_barkod = basliklar.get("Barkod")
col_eski_fiyat = basliklar.get("Fiyat") 

# Eski "Fiyat" başlığının adını "Fiyat 1" yapalım
if col_eski_fiyat:
    ws.cell(row=1, column=col_eski_fiyat).value = "Fiyat 1"

# Fiyat 2 için en sağa yeni bir sütun açalım (Resimlerin olduğu sütunu kaydırmamak için en sağa ekliyoruz)
col_fiyat2 = ws.max_column + 1
ws.cell(row=1, column=col_fiyat2).value = "Fiyat 2"

# 3. Satır satır gez ve Fiyat 2'leri eşleştirip ekle
print("Fiyat 2'ler Excel'e işleniyor...")
eklenen = 0
for satir in range(2, ws.max_row + 1):
    hucre_barkod = ws.cell(row=satir, column=col_barkod).value
    if not hucre_barkod:
        continue
        
    for parca in str(hucre_barkod).split():
        anahtar = normalize(parca)
        if anahtar in fiyat2_sozlugu:
            yeni_f2 = fiyat2_sozlugu[anahtar]
            try:
                yeni_f2 = float(yeni_f2)
            except:
                pass
            ws.cell(row=satir, column=col_fiyat2).value = yeni_f2
            eklenen += 1
            break

wb.save(CIKTI_DOSYA)
print("\n" + "=" * 30)
print("--- İŞLEM TAMAMLANDI ---")
print(f"Başarıyla eklenen Fiyat 2 sayısı: {eklenen}")
print(f"Tüm verilerin tam olduğu yeni dosyanız: {CIKTI_DOSYA}")
print("=" * 30)