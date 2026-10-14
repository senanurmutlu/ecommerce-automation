# -*- coding: utf-8 -*-
import time
import base64
import requests
import urllib.parse
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from webdriver_manager.chrome import ChromeDriverManager

# ---------------- AYARLAR ----------------
GIRIS_DOSYASI = "HARUN_CALISMA_KDV_FIYAT_GUNCEL.xlsx"
SAYFA_ADI     = "Ürünlerinizi Burada Listeleyin"
CIKTI_DOSYASI = "HARUN_CALISMA_RESIMLI_SON.xlsx" # Adını değiştirdim ki eskisiyle karışmasın

URUN_SUTUNU   = "Ürün Adı"
RESIM_SUTUNU  = "Görseller"
RESIM_BOYUTU  = 80  
# -----------------------------------------

print("Excel dosyası yükleniyor...")
wb = load_workbook(GIRIS_DOSYASI)
ws = wb[SAYFA_ADI]

basliklar = {str(c.value).strip(): c.column for c in ws[1] if c.value}
col_urun = basliklar.get(URUN_SUTUNU)

col_resim = ws.max_column + 1
ws.cell(row=1, column=col_resim).value = RESIM_SUTUNU
col_resim_harfi = get_column_letter(col_resim)
ws.column_dimensions[col_resim_harfi].width = 12

print("Tarayıcı başlatılıyor...")

options = Options()
options.add_argument("--window-size=800,600")
options.add_argument("--disable-blink-features=AutomationControlled") 
options.add_experimental_option("excludeSwitches", ["enable-automation"])
options.add_experimental_option('useAutomationExtension', False)
# Bu ayar sayfanın tamamen yüklenmesini beklemez, işi çok hızlandırır!
options.page_load_strategy = 'eager' 

driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)

bulunan, bulunamayan, resim_hatasi = 0, 0, 0

for satir in range(2, ws.max_row + 1):
    raw_urun = ws.cell(row=satir, column=col_urun).value
    if not raw_urun:
        continue
        
    urun_adi = str(raw_urun).strip()
    if not urun_adi:
        continue
        
    try:
        url = f"https://www.google.com/search?tbm=isch&q={urllib.parse.quote(urun_adi)}"
        driver.get(url)
        time.sleep(1) # Bekleme süresini de 1 saniyeye düşürdüm
        
        resimler = driver.find_elements(By.TAG_NAME, "img")
        img_src = None
        
        for r in resimler:
            try:
                # İŞTE BİZİ KURTARACAK OLAN SATIR: 
                # Sadece genişliği 60 pikselden büyük olan resimleri kabul et (Klavye, logo vb. çöpe!)
                if r.size['width'] > 60 and r.size['height'] > 60:
                    src = r.get_attribute("src")
                    if src and ("data:image" in src or "http" in src):
                        img_src = src
                        break # Gerçek ürünü bulduk, döngüden çık
            except:
                continue
                
        if img_src:
            try:
                if img_src.startswith("data:image"):
                    base64_data = img_src.split(",")[1]
                    img_data = base64.b64decode(base64_data)
                else:
                    resp = requests.get(img_src, timeout=5)
                    img_data = resp.content
                    
                pil_img = PILImage.open(BytesIO(img_data))
                if pil_img.mode != "RGB":
                    pil_img = pil_img.convert("RGB")
                    
                temiz_img_io = BytesIO()
                pil_img.save(temiz_img_io, format="PNG")
                temiz_img_io.seek(0)
                
                img = OpenpyxlImage(temiz_img_io)
                img.width = RESIM_BOYUTU
                img.height = RESIM_BOYUTU
                
                ws.add_image(img, f"{col_resim_harfi}{satir}")
                ws.row_dimensions[satir].height = RESIM_BOYUTU * 0.75
                
                print(f"[+] Eklendi: {urun_adi}")
                bulunan += 1
            except Exception:
                print(f"[!] İşlenemedi: {urun_adi}")
                resim_hatasi += 1
        else:
            print(f"[-] Bulunamadı: {urun_adi}")
            bulunamayan += 1
            
    except Exception:
        print(f"[!] Tarayıcı hatası ({urun_adi})")
        resim_hatasi += 1

driver.quit()
wb.save(CIKTI_DOSYASI)
print("\n--- BÜYÜK GEÇMİŞ OLSUN, İŞLEM TAMAMLANDI! ---")